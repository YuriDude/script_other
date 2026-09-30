import os
import time
import requests
from openpyxl import load_workbook

# ============================================================
# Configuration
# ============================================================

# Excel file to process
INPUT_FILE = "TV1.xlsx"

# Second column = column B
URL_COLUMN = 2

# Headers to add after the existing last column
RESPONSE_HEADER = "URL_Response"
NEW_SITE_HEADER = "NEW_Site"

# Request settings
TIMEOUT = 15

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/140.0 Safari/537.36"
    )
}


# ============================================================
# Extract URL from Excel hyperlink
# ============================================================

def get_url_from_cell(cell):
    """
    Returns the actual hyperlink URL stored in the Excel cell.

    Example:
        Visible text: Microsoft Website
        Actual URL:   https://www.microsoft.com/
    """

    if cell.hyperlink:
        return cell.hyperlink.target

    return None


# ============================================================
# Validate URL
# ============================================================

def validate_url(url):
    """
    Requests the URL and returns:

        response_code
        final_url

    requests automatically follows redirects.
    """

    if not url:
        return None, None

    try:
        response = requests.get(
            url,
            headers=HEADERS,
            timeout=TIMEOUT,
            allow_redirects=True
        )

        return response.status_code, response.url

    except requests.exceptions.Timeout:
        return "TIMEOUT", None

    except requests.exceptions.TooManyRedirects:
        return "TOO_MANY_REDIRECTS", None

    except requests.exceptions.SSLError:
        return "SSL_ERROR", None

    except requests.exceptions.ConnectionError:
        return "CONNECTION_ERROR", None

    except requests.exceptions.RequestException as e:
        return f"ERROR: {type(e).__name__}", None


# ============================================================
# Main processing
# ============================================================

def process_excel():

    if not os.path.exists(INPUT_FILE):
        print(f"ERROR: File not found: {INPUT_FILE}")
        return

    print(f"Opening: {INPUT_FILE}")

    # Load the original workbook
    workbook = load_workbook(INPUT_FILE)

    for worksheet in workbook.worksheets:

        print(f"\nProcessing worksheet: {worksheet.title}")

        # Find the existing last column
        last_column = worksheet.max_column

        response_column = last_column + 1
        new_site_column = last_column + 2

        # Add headers
        worksheet.cell(
            row=1,
            column=response_column,
            value=RESPONSE_HEADER
        )

        worksheet.cell(
            row=1,
            column=new_site_column,
            value=NEW_SITE_HEADER
        )

        # Process every data row
        for row in range(2, worksheet.max_row + 1):

            cell = worksheet.cell(
                row=row,
                column=URL_COLUMN
            )

            # Get the hidden hyperlink
            url = get_url_from_cell(cell)

            # No hyperlink in this cell
            if not url:
                continue

            print(f"Row {row}: {url}")

            response_code, final_url = validate_url(url)

            # Write response code
            worksheet.cell(
                row=row,
                column=response_column,
                value=response_code
            )

            # Write redirected URL only if it changed
            if final_url and final_url.rstrip("/") != url.rstrip("/"):
                worksheet.cell(
                    row=row,
                    column=new_site_column,
                    value=final_url
                )

                print(
                    f"  -> {response_code} "
                    f"REDIRECTED TO: {final_url}"
                )

            else:
                print(f"  -> {response_code}")

            # Small delay between requests
            time.sleep(0.1)

    # ========================================================
    # IMPORTANT:
    # Save back to the SAME file
    # ========================================================

    workbook.save(INPUT_FILE)

    print("\n===================================")
    print("Processing complete")
    print(f"Updated file: {INPUT_FILE}")
    print("===================================")


# ============================================================
# Run
# ============================================================

if __name__ == "__main__":
    process_excel()
